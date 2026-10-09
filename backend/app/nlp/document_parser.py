import os
import re
from pathlib import Path
from typing import Dict, Any, Tuple, Optional
import pymupdf  # PyMuPDF modern import
import docx

from app.nlp.preprocessor import clean_text

SECTION_KEYWORDS = {
    "summary": ["summary", "professional summary", "profile", "about me", "objective", "career objective"],
    "experience": ["experience", "work experience", "employment history", "professional experience", "internships", "work history"],
    "education": ["education", "academic background", "qualifications", "educational background", "academic qualifications"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "skill set", "technical expertise", "tools and technologies"],
    "projects": ["projects", "academic projects", "key projects", "personal projects", "portfolio"],
    "certifications": ["certifications", "certificates", "courses", "licenses", "achievements", "accomplishments"]
}

def extract_text_from_pdf(file_path: str) -> str:
    text_parts = []
    try:
        doc = pymupdf.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text_parts.append(page.get_text())
        doc.close()
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")
    
    return clean_text("\n".join(text_parts))

def extract_text_from_docx(file_path: str) -> str:
    text_parts = []
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
        
    return clean_text("\n".join(text_parts))

def extract_text_from_file(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()
    
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return clean_text(f.read())
    else:
        raise ValueError(f"Unsupported file format: {ext}. Please upload a PDF or DOCX file.")

def segment_resume_sections(text: str) -> Dict[str, str]:
    sections: Dict[str, list] = {
        "header": [],
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "projects": [],
        "certifications": [],
        "other": []
    }
    
    lines = text.split("\n")
    current_section = "header"
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        
        normalized_header = re.sub(r"[^a-zA-Z\s]", "", stripped).strip().lower()
        matched_section = None
        
        for sec_name, keywords in SECTION_KEYWORDS.items():
            if normalized_header in keywords or any(normalized_header == kw for kw in keywords):
                matched_section = sec_name
                break
            elif len(normalized_header.split()) <= 4 and any(kw in normalized_header for kw in keywords):
                matched_section = sec_name
                break
        
        if matched_section and len(stripped) < 45:
            current_section = matched_section
        else:
            sections[current_section].append(stripped)
            
    return {sec: "\n".join(content_lines) for sec, content_lines in sections.items()}
